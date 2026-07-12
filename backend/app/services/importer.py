from dataclasses import dataclass
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.enums import GameSource, GameStatus, RuleStatus, SeasonStatus
from app.models.game import Game, GamePlayer
from app.models.player import Player, PlayerCharacter
from app.models.rating import SeasonPlayerClassStats, SeasonPlayerStats
from app.models.scoring import ScoringRuleVersion
from app.models.season import Season
from app.services.names import normalize_player_name

IMPORTANT_SHEETS = {
    "전적확인",
    "전적이력",
    "시즌목록",
    "듀오전적",
    "듀오전적전체",
    "포인트변동",
    "점수MMR",
    "점수MMR_디스코드",
    "사용자목록",
    "기록입력_디스코드",
}

CLASS_COLUMNS = [
    ("드루", 1, 2),
    ("어쎄", 3, 4),
    ("네크", 5, 6),
    ("슴딘", 7, 8),
]

GAME_CLASS_COLUMNS = [
    ("드루", 9, 13),
    ("어쎄", 10, 14),
    ("네크", 11, 15),
    ("슴딘", 12, 16),
]


@dataclass(frozen=True)
class SheetPreview:
    name: str
    max_row: int
    max_column: int
    header_row: int | None
    headers: list[str]


@dataclass(frozen=True)
class WorkbookSummary:
    path: str
    sheet_count: int
    sheets: list[SheetPreview]


@dataclass(frozen=True)
class LegacyImportResult:
    seasons: int = 0
    players: int = 0
    player_characters: int = 0
    season_stats: int = 0
    class_stats: int = 0
    games: int = 0
    game_players: int = 0


def inspect_workbook(path: str | Path) -> WorkbookSummary:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries
    except ImportError as exc:
        raise RuntimeError("openpyxl is required. Install requirements-dev.txt first.") from exc

    workbook_path = Path(path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    previews: list[SheetPreview] = []

    for worksheet in workbook.worksheets:
        if worksheet.title not in IMPORTANT_SHEETS and not worksheet.title.startswith(
            "전적확인_시즌"
        ):
            continue
        header_row, headers = detect_header(worksheet)
        max_row, max_column = sheet_dimensions(worksheet, range_boundaries)
        previews.append(
            SheetPreview(
                name=worksheet.title,
                max_row=max_row,
                max_column=max_column,
                header_row=header_row,
                headers=headers,
            )
        )

    return WorkbookSummary(
        path=str(workbook_path),
        sheet_count=len(workbook.sheetnames),
        sheets=previews,
    )


def import_legacy_workbook(path: str | Path, db: Session) -> LegacyImportResult:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required. Install requirements-dev.txt first.") from exc

    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    importer = LegacyWorkbookImporter(workbook=workbook, db=db)
    return importer.run()


class LegacyWorkbookImporter:
    def __init__(self, workbook: Any, db: Session) -> None:
        self.workbook = workbook
        self.db = db
        self.current_season_name: str | None = None
        self.counts = {
            "seasons": 0,
            "players": 0,
            "player_characters": 0,
            "season_stats": 0,
            "class_stats": 0,
            "games": 0,
            "game_players": 0,
        }

    def run(self) -> LegacyImportResult:
        rule = self.ensure_legacy_rule()
        self.import_seasons(rule)
        self.import_players()
        self.import_season_snapshots()
        self.import_games()
        self.db.commit()
        return LegacyImportResult(**self.counts)

    def ensure_legacy_rule(self) -> ScoringRuleVersion:
        rule = self.db.scalar(
            select(ScoringRuleVersion).where(ScoringRuleVersion.name == "Legacy Sheet Snapshot")
        )
        if rule is not None:
            return rule

        rule = ScoringRuleVersion(
            name="Legacy Sheet Snapshot",
            status=RuleStatus.RETIRED.value,
            config={"source": "google_sheet", "recalculation": False},
            notes="Imported closed-season snapshot. Not used for new seasons.",
        )
        self.db.add(rule)
        self.db.flush()
        return rule

    def import_seasons(self, rule: ScoringRuleVersion) -> None:
        worksheet = self.workbook["시즌목록"]
        for row_index, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            if row_index == 2:
                current_name = clean_text(cell(row, 4))
                if current_name:
                    self.current_season_name = current_name
                    self.upsert_season(
                        name=current_name,
                        starts_at=parse_datetime(cell(row, 5)),
                        ends_at=parse_datetime(cell(row, 6)),
                        rule_id=rule.id,
                    )

            historical_name = clean_text(cell(row, 0))
            if not historical_name:
                continue

            self.upsert_season(
                name=historical_name,
                starts_at=parse_datetime(cell(row, 1)),
                ends_at=parse_datetime(cell(row, 2)),
                rule_id=rule.id,
            )

    def upsert_season(
        self,
        name: str,
        starts_at: datetime | None,
        ends_at: datetime | None,
        rule_id: int,
    ) -> Season:
        season = self.db.scalar(select(Season).where(Season.name == name))
        if season is None:
            season = Season(name=name, sort_order=season_sort_order(name))
            self.db.add(season)
            self.counts["seasons"] += 1

        season.status = SeasonStatus.CLOSED.value
        season.starts_at = starts_at
        season.ends_at = ends_at
        season.scoring_rule_version_id = rule_id
        season.is_imported_snapshot = True
        season.legacy_sheet_name = (
            "전적확인" if name == self.current_season_name else f"전적확인_{name}"
        )
        return season

    def import_players(self) -> None:
        worksheet = self.workbook["사용자목록"]
        class_map = [("드루", 1), ("어쎄", 2), ("네크", 3), ("슴딘", 4)]
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            player_name = clean_text(cell(row, 0))
            if not is_valid_player_name(player_name):
                continue

            player, created = self.get_or_create_player(player_name)
            if created:
                self.counts["players"] += 1
            player.current_tier = clean_text(cell(row, 14)) or None

            for class_name, column_index in class_map:
                character_name = clean_text(cell(row, column_index))
                if not character_name:
                    continue
                if self.upsert_player_character(player, class_name, character_name):
                    self.counts["player_characters"] += 1

    def import_season_snapshots(self) -> None:
        for sheet_name in self.workbook.sheetnames:
            season_name = self.season_name_for_rank_sheet(sheet_name)
            if season_name is None:
                continue
            season = self.db.scalar(select(Season).where(Season.name == season_name))
            if season is None:
                continue
            self.import_rank_sheet(self.workbook[sheet_name], season)

    def season_name_for_rank_sheet(self, sheet_name: str) -> str | None:
        if sheet_name == "전적확인":
            return self.current_season_name
        if sheet_name.startswith("전적확인_"):
            return sheet_name.removeprefix("전적확인_")
        return None

    def import_rank_sheet(self, worksheet: Any, season: Season) -> None:
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            player_name = clean_text(cell(row, 0))
            if not is_valid_player_name(player_name):
                continue

            wins = parse_int(cell(row, 9), default=0)
            losses = parse_int(cell(row, 10), default=0)
            rating = parse_int(cell(row, 17), default=0)
            rank = parse_int(cell(row, 16), default=None)
            if wins == 0 and losses == 0 and rating == 0:
                continue

            player, created = self.get_or_create_player(player_name)
            if created:
                self.counts["players"] += 1

            stats = self.db.scalar(
                select(SeasonPlayerStats).where(
                    SeasonPlayerStats.season_id == season.id,
                    SeasonPlayerStats.player_id == player.id,
                )
            )
            if stats is None:
                stats = SeasonPlayerStats(season_id=season.id, player_id=player.id, rating=0)
                self.db.add(stats)
                self.counts["season_stats"] += 1

            stats.rank = rank
            stats.rating = rating
            stats.games_played = wins + losses
            stats.wins = wins
            stats.losses = losses
            stats.win_rate = parse_float(cell(row, 11))
            stats.tier = clean_text(cell(row, 21)) or None
            stats.is_snapshot = True

            for class_name, win_column, loss_column in CLASS_COLUMNS:
                class_wins = parse_int(cell(row, win_column), default=0)
                class_losses = parse_int(cell(row, loss_column), default=0)
                if class_wins == 0 and class_losses == 0:
                    continue
                if self.upsert_class_stats(season, player, class_name, class_wins, class_losses):
                    self.counts["class_stats"] += 1

    def import_games(self) -> None:
        worksheet = self.workbook["전적이력"]
        for row_index, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            legacy_game_id = clean_text(cell(row, 0))
            if not legacy_game_id or not legacy_game_id.startswith("G"):
                continue

            existing_game = self.db.scalar(
                select(Game).where(
                    Game.source == GameSource.LEGACY_IMPORT.value,
                    Game.legacy_row_number == row_index,
                )
            )
            if existing_game is not None:
                continue

            season_name = clean_text(cell(row, 7))
            season = self.db.scalar(select(Season).where(Season.name == season_name))
            played_at = parse_datetime(cell(row, 1))
            score_a = parse_int(cell(row, 3), default=None)
            score_b = parse_int(cell(row, 5), default=None)
            if season is None or played_at is None or score_a is None or score_b is None:
                continue

            game = Game(
                id=str(uuid4()),
                season_id=season.id,
                legacy_game_id=legacy_game_id,
                legacy_row_number=row_index,
                played_at=played_at,
                source=GameSource.LEGACY_IMPORT.value,
                status=GameStatus.ACTIVE.value,
                winner_side="A",
                score_a=score_a,
                score_b=score_b,
                scoring_rule_version_id=season.scoring_rule_version_id,
            )
            self.db.add(game)
            self.db.flush()
            self.counts["games"] += 1

            self.import_game_players(game, row)

    def import_game_players(self, game: Game, row: tuple[Any, ...]) -> None:
        all_character_players = {
            item.strip()
            for item in clean_text(cell(row, 8)).split(",")
            if item and item.strip()
        }
        game_class_columns = enumerate(GAME_CLASS_COLUMNS, start=1)
        for slot, (class_name, winner_column, loser_column) in game_class_columns:
            winner_name = clean_text(cell(row, winner_column))
            if is_valid_player_name(winner_name):
                player, created = self.get_or_create_player(winner_name)
                if created:
                    self.counts["players"] += 1
                self.db.add(
                    GamePlayer(
                        game_id=game.id,
                        player_id=player.id,
                        side="A",
                        slot=slot,
                        class_name_at_game=class_name,
                        is_all_character_bonus=winner_name in all_character_players,
                    )
                )
                self.counts["game_players"] += 1

            loser_name = clean_text(cell(row, loser_column))
            if is_valid_player_name(loser_name):
                player, created = self.get_or_create_player(loser_name)
                if created:
                    self.counts["players"] += 1
                self.db.add(
                    GamePlayer(
                        game_id=game.id,
                        player_id=player.id,
                        side="B",
                        slot=slot,
                        class_name_at_game=class_name,
                        is_all_character_bonus=loser_name in all_character_players,
                    )
                )
                self.counts["game_players"] += 1

    def upsert_class_stats(
        self,
        season: Season,
        player: Player,
        class_name: str,
        wins: int,
        losses: int,
    ) -> bool:
        class_stats = self.db.scalar(
            select(SeasonPlayerClassStats).where(
                SeasonPlayerClassStats.season_id == season.id,
                SeasonPlayerClassStats.player_id == player.id,
                SeasonPlayerClassStats.class_name == class_name,
            )
        )
        created = False
        if class_stats is None:
            class_stats = SeasonPlayerClassStats(
                season_id=season.id,
                player_id=player.id,
                class_name=class_name,
            )
            self.db.add(class_stats)
            created = True

        class_stats.wins = wins
        class_stats.losses = losses
        class_stats.win_rate = round(wins / (wins + losses) * 100, 1) if wins + losses else None
        class_stats.is_snapshot = True
        return created

    def upsert_player_character(
        self,
        player: Player,
        class_name: str,
        character_name: str,
    ) -> bool:
        character = self.db.scalar(
            select(PlayerCharacter).where(
                PlayerCharacter.player_id == player.id,
                PlayerCharacter.class_name == class_name,
            )
        )
        created = False
        if character is None:
            character = PlayerCharacter(player_id=player.id, class_name=class_name)
            self.db.add(character)
            created = True

        character.character_name = character_name
        character.normalized_character_name = normalize_player_name(character_name)
        return created

    def get_or_create_player(self, display_name: str) -> tuple[Player, bool]:
        normalized_name = normalize_player_name(display_name)
        player = self.db.scalar(select(Player).where(Player.normalized_name == normalized_name))
        if player is not None:
            return player, False

        player = Player(
            display_name=display_name.strip(),
            normalized_name=normalized_name,
        )
        self.db.add(player)
        self.db.flush()
        return player, True


def cell(row: tuple[Any, ...], index: int) -> Any:
    return row[index] if index < len(row) else None


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def is_valid_player_name(value: str) -> bool:
    if not value:
        return False
    return not value.startswith("#")


def parse_int(value: Any, default: int | None = 0) -> int | None:
    if value is None or value == "":
        return default
    try:
        return int(float(str(value).strip()))
    except ValueError:
        return default


def parse_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(str(value).strip())
    except ValueError:
        return None


def parse_datetime(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    if isinstance(value, date):
        return datetime.combine(value, time.min, tzinfo=timezone.utc)

    text = str(value).strip()
    for date_format in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d",
        "%y.%m.%d",
    ):
        try:
            parsed = datetime.strptime(text, date_format)
        except ValueError:
            continue
        return parsed.replace(tzinfo=timezone.utc)
    return None


def season_sort_order(name: str) -> int:
    digits = "".join(character for character in name if character.isdigit())
    return int(digits) if digits else 0


def sheet_dimensions(worksheet: Any, range_boundaries: Any) -> tuple[int, int]:
    if worksheet.max_row and worksheet.max_column:
        return worksheet.max_row, worksheet.max_column

    try:
        dimension = worksheet.calculate_dimension(force=True)
    except TypeError:
        dimension = worksheet.calculate_dimension()
    except ValueError:
        return 0, 0

    if not dimension or dimension == "A1:A1":
        return worksheet.max_row or 0, worksheet.max_column or 0

    _, _, max_column, max_row = range_boundaries(dimension)
    return max_row, max_column


def detect_header(worksheet: Any, max_rows: int = 12) -> tuple[int | None, list[str]]:
    bounded_max_row = min(max_rows, worksheet.max_row) if worksheet.max_row else max_rows
    for row_index, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=bounded_max_row, values_only=True),
        start=1,
    ):
        values = [format_cell(value) for value in row]
        non_empty = [value for value in values if value]
        if len(non_empty) >= 2:
            return row_index, values
    return None, []


def format_cell(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if len(text) > 90:
        return f"{text[:87]}..."
    return text
